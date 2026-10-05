"""
Evaluation Sheet Writer - Gera as planilhas de avaliação (uma cópia por avaliador)
"""
from pathlib import Path
from typing import List

import pandas as pd
from loguru import logger
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation


class EvaluationSheetWriter:
    """
    Gera as planilhas cegas de avaliação.

    Responsabilidades:
    - Escrever só id_anonimo e texto (demais colunas em branco)
    - Aplicar validação de lista em rotulo_escolhido (classes + opção de informação
      insuficiente), qual_outro_rotulo (classes) e outro_rotulo_possivel (sim/não)
    - Gravar as cópias dos avaliadores com a mesma ordem de linhas
    - Não sobrescrever planilhas existentes (podem já estar preenchidas)
    """

    COLUMNS = ["id_anonimo", "texto", "rotulo_escolhido", "outro_rotulo_possivel", "qual_outro_rotulo", "observacao"]
    YES_NO = ["sim", "não"]
    EXCEL_CELL_LIMIT = 32_000
    EXCEL_LIST_LIMIT = 255

    def __init__(self, class_names: List[str], insufficient_option: str):
        self.class_names = class_names
        self.chosen_options = class_names + [insufficient_option]
        for options in (self.chosen_options, self.YES_NO):
            if len(",".join(options)) > self.EXCEL_LIST_LIMIT or any("," in o for o in options):
                raise ValueError(f"Lista inválida para validação do Excel: {options}")
        logger.debug(f"EvaluationSheetWriter inicializado ({len(class_names)} classes)")

    def _clean_text(self, text: str) -> str:
        text = ILLEGAL_CHARACTERS_RE.sub("", str(text))
        if len(text) > self.EXCEL_CELL_LIMIT:
            logger.warning(f"Texto truncado em {self.EXCEL_CELL_LIMIT} caracteres (limite da célula do Excel)")
            text = text[: self.EXCEL_CELL_LIMIT] + " [...]"
        return text

    def _list_validation(self, options: List[str], column: str, n_rows: int) -> DataValidation:
        validation = DataValidation(
            type="list",
            formula1=f'"{",".join(options)}"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Valor inválido",
            error="Escolha um valor da lista.",
        )
        validation.add(f"{column}2:{column}{n_rows + 1}")
        return validation

    def build(self, rows: pd.DataFrame) -> Workbook:
        """Workbook a partir de um DataFrame com id_anonimo e texto, já na ordem final."""
        wb = Workbook()
        ws = wb.active
        ws.title = "avaliacao"
        ws.append(self.COLUMNS)

        for id_anonimo, texto in zip(rows["id_anonimo"], rows["texto"]):
            ws.append([id_anonimo, self._clean_text(texto), None, None, None, None])

        n = len(rows)
        ws.add_data_validation(self._list_validation(self.chosen_options, "C", n))
        ws.add_data_validation(self._list_validation(self.YES_NO, "D", n))
        ws.add_data_validation(self._list_validation(self.class_names, "E", n))

        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", (14, 90, 26, 22, 26, 40)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        return wb

    def write_copies(self, rows: pd.DataFrame, evaluators: List[str], output_dir: Path) -> List[Path]:
        """Uma planilha por avaliador, todas a partir do mesmo `rows`."""
        paths = []
        for evaluator in evaluators:
            path = Path(output_dir) / f"planilha_avaliacao_{evaluator}.xlsx"
            if path.exists():
                logger.info(f"Planilha já existe, mantida: {path.name}")
            else:
                self.build(rows).save(path)
                logger.success(f"Planilha salva: {path}")
            paths.append(path)
        return paths
